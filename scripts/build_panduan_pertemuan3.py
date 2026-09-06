#!/usr/bin/env python3
"""Bangun PANDUAN_PERTEMUAN3.docx (buku panduan tugas Pertemuan 3).

Struktur mengikuti pola umum modul/guidebook praktikum: cover, daftar isi,
daftar gambar, Percobaan 1-9 (tujuan, langkah, kode, gambar, output,
pembahasan), rangkuman, lembar kerja, daftar pustaka. Tanpa struktur
BAB ala laporan formal.

Format visual meniru laporan KKN: A4, margin 1 inci, Times New Roman 12
(dari cover sampai isi), judul tengah tebal, isi justify spasi 1,5.
Cuplikan kode dan output ditempatkan dalam tabel 1 x 1 selebar margin
dengan font Courier New 10. Slot screenshot tabel 1 x 1 selebar margin,
diisi otomatis dari berkas assets/colab/pertemuan3_P{n}.png bila ada.

Pakai: uv run --python 3.13 --with python-docx python3 scripts/build_panduan_pertemuan3.py
"""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).resolve().parents[1] / "PANDUAN_PERTEMUAN3.docx"
LOGO = Path(__file__).resolve().parents[1] / "assets" / "logo-unud.png"
COLAB = Path(__file__).resolve().parents[1] / "assets" / "colab"
TNR = "Times New Roman"
BLACK = RGBColor(0, 0, 0)
USABLE_CM = 21.0 - 2.54 - 2.54  # lebar area isi A4 margin 1 inci

ANGGOTA = [
    ("1", "2305551036", "Deliana Br Manalu"),
    ("2", "2305551076", "Ravi Arnan Irianto"),
    ("3", "2305551144", "Ezza Putra Wibawa"),
    ("4", "2305551173", "Devin"),
]


def fit_margin(tbl):
    """Buat tabel selebar area isi (margin 1 inci), lebar dibagi rata per kolom."""
    ncol = max(len(r.cells) for r in tbl.rows) if tbl.rows else 1
    col_w = Cm(USABLE_CM / ncol)
    tbl.autofit = False
    tbl.allow_autofit = False
    for row in tbl.rows:
        for c in row.cells:
            c.width = col_w
    tblPr = tbl._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    twips = int(USABLE_CM * 567 / ncol)
    for gridCol in tbl._tbl.tblGrid.findall(qn("w:gridCol")):
        gridCol.set(qn("w:w"), str(twips))
    return tbl


def body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.first_line_indent = Cm(1.0)
    p.paragraph_format.space_after = Pt(6)
    if bold_prefix:
        r = p.add_run(bold_prefix)
        r.font.name = TNR
        r.font.size = Pt(12)
        r.bold = True
    r = p.add_run(text)
    r.font.name = TNR
    r.font.size = Pt(12)
    return p


def heading1(doc, text):
    p = doc.add_paragraph()
    p.style = doc.styles["Heading 1"]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text.upper())
    r.font.name = TNR
    r.font.size = Pt(12)
    r.bold = True
    r.font.color.rgb = BLACK
    return p


def heading2(doc, text):
    p = doc.add_paragraph()
    p.style = doc.styles["Heading 2"]
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.font.name = TNR
    r.font.size = Pt(12)
    r.bold = True
    r.font.color.rgb = BLACK
    return p


def numbered(doc, items):
    for i, langkah in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(f"{i}. {langkah}")
        r.font.name = TNR
        r.font.size = Pt(12)


def caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.font.name = TNR
    r.font.size = Pt(12)
    r.bold = True
    return p


def code_block(doc, lines):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.style = "Table Grid"
    fit_margin(tbl)
    cell = tbl.cell(0, 0)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.0
    for i, line in enumerate(lines):
        r = p.add_run(line if line else " ")
        r.font.name = "Courier New"
        r.font.size = Pt(10)
        if i < len(lines) - 1:
            r.add_break()
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return tbl


def output_block(doc, lines):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.style = "Table Grid"
    fit_margin(tbl)
    cell = tbl.cell(0, 0)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.0
    for i, line in enumerate(lines):
        r = p.add_run(line if line else " ")
        r.font.name = "Courier New"
        r.font.size = Pt(10)
        if i < len(lines) - 1:
            r.add_break()
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return tbl


def slot_screenshot(doc, gambar_no, keterangan):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.style = "Table Grid"
    fit_margin(tbl)
    cell = tbl.cell(0, 0)
    cell.width = Cm(USABLE_CM)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    png = COLAB / f"pertemuan3_P{gambar_no}.png"
    if png.exists():
        r = p.add_run()
        r.add_picture(str(png), width=Cm(USABLE_CM))
    else:
        r = p.add_run(f"[TEMPEL SCREENSHOT COLAB DI SINI: {keterangan}]")
        r.font.name = TNR
        r.font.size = Pt(11)
        r.italic = True
    caption(doc, f"Gambar {gambar_no}. {keterangan}")


def centered(doc, text, bold=False, size=12, space_after=6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text)
    r.font.name = TNR
    r.font.size = Pt(size)
    r.bold = bold
    return p


def toc_entry(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = TNR
    r.font.size = Pt(12)
    return p


def result_table(doc, headers, rows):
    tbl = doc.add_table(rows=1 + len(rows), cols=len(headers))
    tbl.style = "Table Grid"
    fit_margin(tbl)
    for j, h in enumerate(headers):
        c = tbl.cell(0, j)
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.font.name = TNR
        r.font.size = Pt(12)
        r.bold = True
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            c = tbl.cell(i, j)
            c.text = ""
            r = c.paragraphs[0].add_run(str(val))
            r.font.name = TNR
            r.font.size = Pt(12)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)


def percobaan(doc, no, judul, tujuan, langkah, kode, gambar_ket, output, bahasan):
    doc.add_page_break()
    heading1(doc, f"Percobaan {no}: {judul}")
    heading2(doc, "Tujuan")
    body(doc, tujuan)
    heading2(doc, "Langkah percobaan")
    numbered(doc, langkah)
    heading2(doc, "Kode program")
    code_block(doc, kode)
    heading2(doc, "Hasil (screenshot)")
    slot_screenshot(doc, no, gambar_ket)
    heading2(doc, "Output")
    output_block(doc, output)
    heading2(doc, "Pembahasan")
    for teks in bahasan:
        body(doc, teks)


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Cm(2.54)

    normal = doc.styles["Normal"]
    normal.font.name = TNR
    normal.font.size = Pt(12)

    # ---- COVER (ala laporan KK dampingan: judul, logo, disusun oleh, institusi, tahun) ----
    if not LOGO.exists():
        raise SystemExit(f"logo tidak ditemukan: {LOGO}")
    centered(doc, "BUKU PANDUAN", bold=True)
    centered(doc, "TUGAS PERTEMUAN 3: PENGKONDISIAN DAN PERULANGAN", bold=True)
    centered(doc, "MATA KULIAH MACHINE LEARNING", bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(12)
    p.add_run().add_picture(str(LOGO), width=Cm(4.1))
    centered(doc, "DISUSUN OLEH:", bold=True)
    centered(doc, "Kelompok 3", bold=True)
    for i, (no, nim, nama) in enumerate(ANGGOTA, start=1):
        centered(doc, f"{i}. {nama} ({nim})", space_after=0)
    doc.add_paragraph()
    centered(doc, "Dosen Pengampu: Bapak Adi Purnawan")
    doc.add_paragraph()
    centered(doc, "PROGRAM STUDI S1 TEKNOLOGI INFORMASI", bold=True)
    centered(doc, "UNIVERSITAS UDAYANA", bold=True)
    centered(doc, "TAHUN 2026", bold=True)

    doc.add_page_break()

    # ---- DAFTAR ISI (tanpa nomor halaman) ----
    heading1(doc, "Daftar Isi")
    for entri in [
        "Daftar Isi",
        "Daftar Gambar",
        "Percobaan 1: Memuat Data",
        "Percobaan 2: Ekspresi Boolean",
        "Percobaan 3: Percabangan If-Elif-Else",
        "Percobaan 4: Perulangan For",
        "Percobaan 5: Perulangan While",
        "Percobaan 6: Break dan Continue",
        "Percobaan 7: Nested Loop",
        "Percobaan 8: Comprehension",
        "Percobaan 9: Menyaring Data Latih",
        "Rangkuman",
        "Lembar Kerja",
        "Daftar Pustaka",
    ]:
        toc_entry(doc, entri)

    doc.add_page_break()

    # ---- DAFTAR GAMBAR ----
    heading1(doc, "Daftar Gambar")
    for entri in [
        "Gambar 1. Memuat data dan tiga baris pertama",
        "Gambar 2. Evaluasi ekspresi Boolean pada dataset",
        "Gambar 3. Uji batas kategori glukosa dan risiko lima pasien pertama",
        "Gambar 4. Perulangan for dengan enumerate, range, dan zip",
        "Gambar 5. Perulangan while epoch-loss dan pemindaian bmi valid",
        "Gambar 6. Demonstrasi break dan continue",
        "Gambar 7. Nested loop kelompok usia dan hipertensi",
        "Gambar 8. Comprehension penyaringan dan penskalaan bmi",
        "Gambar 9. Penyaringan data latih dan rekapitulasi alasan penolakan",
    ]:
        toc_entry(doc, entri)

    # ---- PERCOBAAN 1 ----
    percobaan(
        doc,
        1,
        "Memuat Data",
        "Menampilkan ukuran dataset dan memeriksa tiga baris pertama sebelum diolah.",
        [
            "Jalankan sel muat data di Colab.",
            "Catat jumlah baris dan kolom yang tercetak.",
            "Perhatikan tiga baris pertama yang ditampilkan.",
        ],
        [
            "import pandas as pd",
            "",
            'URL = "https://raw.githubusercontent.com/ray-project/raydp/master/tutorials/dataset/healthcare-dataset-stroke-data.csv"',
            "df = pd.read_csv(URL)",
            "",
            'print("Jumlah baris :", df.shape[0])',
            'print("Jumlah kolom :", df.shape[1])',
            "df.head(3)",
        ],
        "memuat data dan tiga baris pertama",
        ["Jumlah baris : 5110", "Jumlah kolom : 12"],
        [
            "Data berhasil dimuat sebanyak 5.110 baris dan 12 kolom. Tiga baris pertama seluruhnya berstatus stroke, yang wajar karena dataset ini terurut dengan kasus stroke di bagian awal."
        ],
    )

    # ---- PERCOBAAN 2 ----
    percobaan(
        doc,
        2,
        "Ekspresi Boolean",
        "Mengajukan pertanyaan True/False ke data dengan operator perbandingan, logika, keanggotaan, dan pemeriksaan nilai kosong.",
        [
            "Jalankan contoh Boolean sederhana dari slide.",
            "Jalankan pemeriksaan bmi kosong, Unknown, dan BMI ekstrem.",
            "Jalankan filter gabungan lansia plus hipertensi dengan smoking diketahui.",
        ],
        [
            "akurasi = 0.87",
            "print(akurasi >= 0.80)  # True",
            "",
            'label = "batik"',
            'print(label in ["batik", "endek"])  # True',
            "",
            "nilai = None",
            "print(nilai is None)  # True",
            "",
            'print("bmi kosong       :", int(df["bmi"].isna().sum()))',
            'print("smoking Unknown  :", int((df["smoking_status"] == "Unknown").sum()))',
            'print("bmi di atas 60   :", int((df["bmi"] > 60).sum()))',
            "",
            'mask = (df["age"] >= 60) & (df["hypertension"] == 1) & (df["smoking_status"] != "Unknown")',
            'print("lansia + hipertensi + smoking diketahui:", int(mask.sum()))',
        ],
        "evaluasi ekspresi Boolean pada dataset",
        [
            "True",
            "True",
            "True",
            "bmi kosong       : 201",
            "smoking Unknown  : 1544",
            "bmi di atas 60   : 13",
            "lansia + hipertensi + smoking diketahui: 261",
        ],
        [
            'Nilai "Unknown" pada smoking_status menempati sekitar 30 persen data. Itu adalah missing value tersembunyi: bukan kategori bermakna sehingga pada Percobaan 9 baris tersebut dilewati dengan continue. Baris bmi kosong hanya sekitar 4 persen dan BMI di atas 60 hanya 13 baris.'
        ],
    )

    # ---- PERCOBAAN 3 ----
    percobaan(
        doc,
        3,
        "Percabangan If-Elif-Else",
        "Menyusun keputusan berlapis dengan urutan batas dari yang paling ketat dan menguji tepat di bawah, pada, dan di atas batas.",
        [
            "Definisikan fungsi kategori_glukosa dan kategori_risiko.",
            "Uji enam nilai batas glukosa.",
            "Terapkan kedua fungsi pada lima baris pertama dataset.",
        ],
        [
            "def kategori_glukosa(g):",
            "    if g >= 200:      # paling ketat dulu",
            '        return "sangat tinggi"',
            "    elif g >= 140:",
            '        return "tinggi"',
            "    elif g >= 100:",
            '        return "waspada"',
            "    else:",
            '        return "normal"',
            "",
            "",
            "def kategori_risiko(age, hyp, heart):",
            "    if heart == 1 and hyp == 1:",
            '        return "tinggi"',
            "    elif heart == 1 or hyp == 1:",
            '        return "sedang"',
            "    elif age >= 60:",
            '        return "waspada"',
            "    else:",
            '        return "rendah"',
            "",
            "",
            "for g in [99.9, 100.0, 139.9, 140.0, 199.9, 200.0]:",
            '    print(g, "->", kategori_glukosa(g))',
        ],
        "uji batas kategori glukosa dan risiko lima pasien pertama",
        [
            "99.9 -> normal",
            "100.0 -> waspada",
            "139.9 -> waspada",
            "140.0 -> tinggi",
            "199.9 -> tinggi",
            "200.0 -> sangat tinggi",
            "",
            "baris 0: glukosa=228.7 (sangat tinggi) risiko=sedang",
            "baris 1: glukosa=202.2 (sangat tinggi) risiko=waspada",
            "baris 2: glukosa=105.9 (waspada) risiko=sedang",
            "baris 3: glukosa=171.2 (tinggi) risiko=rendah",
            "baris 4: glukosa=174.1 (tinggi) risiko=sedang",
        ],
        [
            "Urutan batas menentukan hasil. Jika >= 100 dicek paling dulu maka nilai 250 ikut masuk cabang waspada dan cabang sangat tinggi tidak pernah tercapai. Fungsi kategori_risiko menunjukkan beda and (keduanya harus ada) dengan or (salah satu cukup)."
        ],
    )

    # ---- PERCOBAAN 4 ----
    percobaan(
        doc,
        4,
        "Perulangan For",
        "Menelusuri koleksi yang sudah tersedia item atau jumlah iterasinya dengan enumerate, range, dan zip.",
        [
            "Jalankan contoh enumerate dan range dari slide.",
            "Pasangkan kolom usia dan bmi lima pasien pertama dengan zip.",
            "Hitung rata-rata manual dengan for lalu adu dengan mean() pandas.",
        ],
        [
            "for indeks, skor in enumerate([78, 85, 91]):",
            "    print(indeks, skor)",
            "",
            "for epoch in range(1, 6):",
            '    print("epoch", epoch)',
            "",
            'usia = df["age"].head(5).tolist()',
            'bmi = df["bmi"].head(5).tolist()',
            "for i, (u, b) in enumerate(zip(usia, bmi)):",
            '    print(f"pasien {i}: usia={u} bmi={b}")',
            "",
            "total, n = 0.0, 0",
            'for v in df["age"].head(100):',
            "    total += v",
            "    n += 1",
            "rata_loop = total / n",
            'rata_pandas = df["age"].head(100).mean()',
            'print("rata-rata usia 100 baris (loop):", round(rata_loop, 4))',
            "assert abs(rata_loop - rata_pandas) < 1e-9",
        ],
        "perulangan for dengan enumerate, range, dan zip",
        [
            "0 78",
            "1 85",
            "2 91",
            "",
            "epoch 1",
            "epoch 2",
            "epoch 3",
            "epoch 4",
            "epoch 5",
            "",
            "pasien 0: usia=67.0 bmi=36.6",
            "pasien 1: usia=61.0 bmi=nan",
            "pasien 2: usia=80.0 bmi=32.5",
            "pasien 3: usia=49.0 bmi=34.4",
            "pasien 4: usia=79.0 bmi=24.0",
            "",
            "rata-rata usia 100 baris (loop): 67.28",
        ],
        [
            "Agregasi manual cocok dengan mean() pandas sehingga loop terbukti benar. Rata-rata 67,28 hanya berlaku untuk 100 baris pertama yang semuanya kasus stroke, bukan rata-rata global 43,23. Untuk agregasi kolom penuh, pandas tetap dipakai karena jauh lebih cepat."
        ],
    )

    # ---- PERCOBAAN 5 ----
    percobaan(
        doc,
        5,
        "Perulangan While",
        "Mengulang sampai kondisi terpenuhi dengan variabel pengendali yang selalu diperbarui agar tidak infinite loop.",
        [
            "Jalankan contoh loss menyusut dari slide.",
            "Kumpulkan lima pasien berbmi valid pertama dengan while.",
            "Periksa mengapa loop berhenti di epoch 8.",
        ],
        [
            "epoch, loss = 1, 1.0",
            "while epoch <= 10 and loss > 0.10:",
            "    loss *= 0.72",
            "    print(epoch, round(loss, 3))",
            "    epoch += 1  # tanpa baris ini loop tidak pernah berhenti",
            "",
            "ketemu, i = [], 0",
            "while len(ketemu) < 5 and i < len(df):",
            '    b = df.iloc[i]["bmi"]',
            "    if pd.notna(b):",
            "        ketemu.append((i, b))",
            "    i += 1",
            'print("5 bmi valid pertama (indeks, bmi):", ketemu)',
        ],
        "perulangan while epoch-loss dan pemindaian bmi valid",
        [
            "1 0.72",
            "2 0.518",
            "3 0.373",
            "4 0.269",
            "5 0.193",
            "6 0.139",
            "7 0.1",
            "8 0.072",
            "",
            "5 bmi valid pertama (indeks, bmi): [(0, 36.6), (2, 32.5), (3, 34.4), (4, 24.0), (5, 29.0)]",
        ],
        [
            "Loop berhenti di epoch 8 meski batasnya 10 karena loss sudah di bawah 0,10. Tercetak 8 baris bukan 7 karena nilai epoch 7 sebenarnya sekitar 0,1003 (tercetak 0,1 setelah dibulatkan) sehingga masih di atas 0,10. Inilah beda for dan while: syarat berhenti ditentukan data, bukan jumlah yang tetap."
        ],
    )

    # ---- PERCOBAAN 6 ----
    percobaan(
        doc,
        6,
        "Break dan Continue",
        "Membedakan continue yang melewati satu iterasi dengan break yang keluar dari loop sepenuhnya.",
        [
            "Jalankan contoh slide dengan list berisi None dan negatif.",
            "Pindai 20 bmi pertama dengan aturan lewati NaN dan henti saat ekstrem.",
            "Catat cabang mana yang terpicu dan mana yang tidak.",
        ],
        [
            "for nilai in [12, None, 18, -3, 25]:",
            "    if nilai is None:",
            "        continue",
            "    if nilai < 0:",
            "        break",
            "    print(nilai)",
            "",
            "diproses, dilewati = 0, 0",
            'for b in df["bmi"].head(20).tolist():',
            "    if pd.isna(b):",
            "        dilewati += 1",
            "        continue",
            "    if b > 90:",
            '        print(f"bmi ekstrem {b}, hentikan pindaian demo")',
            "        break",
            "    diproses += 1",
            'print(f"diproses={diproses} dilewati(NaN)={dilewati}")',
        ],
        "demonstrasi break dan continue",
        ["12", "18", "", "diproses=16 dilewati(NaN)=4"],
        [
            "Pada contoh pertama yang tercetak hanya 12 dan 18: None dilewati lalu -3 menghentikan loop sehingga 25 tidak pernah diproses. Pada 20 bmi pertama terdapat 4 nilai kosong sehingga 16 diproses. Cabang break tidak terpicu di sini karena tidak ada bmi di atas 90 pada 20 baris tersebut, dan itu dicatat apa adanya."
        ],
    )

    # ---- PERCOBAAN 7 ----
    percobaan(
        doc,
        7,
        "Nested Loop",
        "Menyilangkan dua dimensi (kelompok usia dan hipertensi) lalu menghitung jumlah pasien dan proporsi stroke tiap sel.",
        [
            "Definisikan fungsi kelompok_usia.",
            "Jalankan nested loop 3 x 2 sel.",
            "Periksa total baris terpetakan sama dengan 5.110.",
        ],
        [
            "def kelompok_usia(a):",
            "    if a < 18:",
            '        return "anak (<18)"',
            "    elif a < 60:",
            '        return "dewasa (18-59)"',
            "    else:",
            '        return "lansia (>=60)"',
            "",
            "",
            "hasil = []",
            'for ku in ["anak (<18)", "dewasa (18-59)", "lansia (>=60)"]:',
            "    for hyp in [0, 1]:",
            '        pot = df[(df["age"].apply(kelompok_usia) == ku) & (df["hypertension"] == hyp)]',
            '        rate = pot["stroke"].mean() if len(pot) else 0.0',
            "        hasil.append((ku, hyp, len(pot), round(float(rate), 4)))",
            '        print(f"{ku} x hipertensi={hyp}: n={len(pot)} stroke_rate={rate:.4f}")',
        ],
        "nested loop kelompok usia dan hipertensi",
        [
            "anak (<18) x hipertensi=0: n=855 stroke_rate=0.0023",
            "anak (<18) x hipertensi=1: n=1 stroke_rate=0.0000",
            "dewasa (18-59) x hipertensi=0: n=2677 stroke_rate=0.0198",
            "dewasa (18-59) x hipertensi=1: n=201 stroke_rate=0.0647",
            "lansia (>=60) x hipertensi=0: n=1080 stroke_rate=0.1185",
            "lansia (>=60) x hipertensi=1: n=296 stroke_rate=0.1791",
            "",
            "total baris terpetakan: 5110",
        ],
        [
            "Keenam sel mempartisi seluruh 5.110 baris tanpa sisa sehingga tiap baris masuk tepat satu sel. Polanya sejalan dengan pengetahuan medis: proporsi stroke naik pada lansia dan pada penderita hipertensi. Nested loop di sini hanya 3 x 2 sel sehingga murah; untuk grid besar sebaiknya diganti groupby."
        ],
    )
    caption(doc, "Tabel 1. Proporsi stroke per kelompok usia dan hipertensi")
    result_table(
        doc,
        ["Kelompok usia", "Hipertensi", "n", "Stroke rate"],
        [
            ["Anak (<18)", "0", "855", "0,0023"],
            ["Anak (<18)", "1", "1", "0,0000"],
            ["Dewasa (18-59)", "0", "2.677", "0,0198"],
            ["Dewasa (18-59)", "1", "201", "0,0647"],
            ["Lansia (>=60)", "0", "1.080", "0,1185"],
            ["Lansia (>=60)", "1", "296", "0,1791"],
        ],
    )

    # ---- PERCOBAAN 8 ----
    percobaan(
        doc,
        8,
        "Comprehension",
        "Menulis transformasi saring dan skala dalam satu baris untuk ekspresi sederhana.",
        [
            "Jalankan contoh slide dengan list berisi None.",
            "Terapkan comprehension saring dan skala pada 10 bmi pertama.",
            "Adu hasilnya dengan dropna().",
        ],
        [
            "data = [10, None, 20, 30]",
            "bersih = [x for x in data if x is not None]",
            "skala = [x / 100 for x in bersih]",
            "print(bersih, skala)",
            "",
            'bmi_10 = df["bmi"].head(10).tolist()',
            "bersih_bmi = [x for x in bmi_10 if pd.notna(x)]",
            "skala_bmi = [x / 100 for x in bersih_bmi]",
            'print("mentah :", bmi_10)',
            'print("bersih :", bersih_bmi)',
            'print("skala  :", [round(v, 4) for v in skala_bmi])',
        ],
        "comprehension penyaringan dan penskalaan bmi",
        [
            "[10, 20, 30] [0.1, 0.2, 0.3]",
            "mentah : [36.6, nan, 32.5, 34.4, 24.0, 29.0, 27.4, 22.8, nan, 24.2]",
            "bersih : [36.6, 32.5, 34.4, 24.0, 29.0, 27.4, 22.8, 24.2]",
            "skala  : [0.366, 0.325, 0.344, 0.24, 0.29, 0.274, 0.228, 0.242]",
        ],
        [
            "Satu baris comprehension menggantikan tiga sampai empat baris loop saring dan hasilnya identik dengan dropna(). Dari 10 nilai terdapat 2 kosong sehingga tersisa 8 nilai bersih. Batasannya keterbacaan: untuk logika berlapis, loop biasa tetap lebih jelas."
        ],
    )

    # ---- PERCOBAAN 9 ----
    percobaan(
        doc,
        9,
        "Menyaring Data Latih",
        "Menerapkan seluruh materi pada kasus nyata: telusuri lalu validasi lalu simpan, lewati fitur kosong, tolak nilai negatif, dan laporkan alasan penolakan.",
        [
            "Jalankan loop penyaringan seluruh 5.110 baris.",
            "Catat jumlah valid, ditolak, dan rincian alasannya.",
            "Periksa valid plus ditolak sama dengan total baris.",
        ],
        [
            "data_valid = []",
            'alasan = {"bmi_kosong": 0, "smoking_unknown": 0, "gender_other": 0, "bmi_negatif": 0}',
            "",
            "for _, baris in df.iterrows():",
            '    if pd.isna(baris["bmi"]):',
            '        alasan["bmi_kosong"] += 1',
            "        continue",
            '    if baris["smoking_status"] == "Unknown":',
            '        alasan["smoking_unknown"] += 1',
            "        continue",
            '    if baris["gender"] == "Other":',
            '        alasan["gender_other"] += 1',
            "        continue",
            '    if baris["bmi"] < 0:',
            '        alasan["bmi_negatif"] += 1',
            "        continue",
            "    data_valid.append(baris)",
            "",
            "n_valid = len(data_valid)",
            "n_tolak = sum(alasan.values())",
            'print(f"valid: {n_valid} ({n_valid/len(df)*100:.1f}%)")',
            'print(f"ditolak: {n_tolak} ({n_tolak/len(df)*100:.1f}%)")',
            "print(alasan)",
        ],
        "penyaringan data latih dan rekapitulasi alasan penolakan",
        [
            "valid: 3425 (67.0%)",
            "ditolak: 1685 (33.0%)",
            "{'bmi_kosong': 201, 'smoking_unknown': 1483, 'gender_other': 1, 'bmi_negatif': 0}",
        ],
        [
            "Penolakan didominasi smoking_status Unknown lalu bmi kosong dan satu baris gender Other. Angka smoking_unknown tercatat 1.483 bukan 1.544 karena 61 baris bermasalah ganda tercatat pada alasan pertama (bmi kosong). Tidak ada bmi negatif sehingga cabang itu menjadi pengaman. Untuk proyek akhir, baris Unknown dan bmi kosong tidak dibuang melainkan ditangani di preprocessing; penyaringan keras di sini adalah demonstrasi kontrol alur, bukan keputusan final pipeline."
        ],
    )
    caption(doc, "Tabel 2. Rekapitulasi penyaringan data latih")
    result_table(
        doc,
        ["Nasib baris", "Jumlah", "Persen"],
        [
            ["Valid", "3.425", "67,0"],
            ["Ditolak (bmi kosong)", "201", "3,9"],
            ["Ditolak (smoking Unknown)", "1.483", "29,0"],
            ["Ditolak (gender Other)", "1", "0,0"],
            ["Ditolak (bmi negatif)", "0", "0,0"],
            ["Total", "5.110", "100,0"],
        ],
    )

    # ---- RANGKUMAN ----
    doc.add_page_break()
    heading1(doc, "Rangkuman")
    body(
        doc,
        "Kontrol alur adalah fondasi pipeline data. Kondisi membuat program adaptif, loop mengotomatisasi proses berulang, dan break serta continue menangani kasus khusus. Kualitas data dimulai dari logika yang jelas.",
    )
    caption(doc, "Tabel 3. Panduan memilih struktur kontrol alur")
    result_table(
        doc,
        ["Kebutuhan", "Struktur yang dipakai"],
        [
            ["Keputusan berlapis", "if-elif-else, batas ketat terlebih dahulu"],
            ["Menelusuri koleksi", "for dengan enumerate atau zip"],
            ["Berhenti ikut kondisi", "while, pengendali selalu maju"],
            ["Mengabaikan satu item", "continue"],
            ["Menghentikan seluruh proses", "break, hanya untuk keadaan darurat"],
        ],
    )

    # ---- LEMBAR KERJA ----
    doc.add_page_break()
    heading1(doc, "Lembar Kerja")
    numbered(
        doc,
        [
            "Mengapa rata-rata usia 100 baris pertama (67,28) jauh di atas rata-rata global (43,23)? Jelaskan dengan susunan dataset.",
            "Mengapa loop while pada Percobaan 5 mencetak 8 baris padahal nilai epoch 7 tercetak 0,1? Jelaskan peran pembulatan.",
            "Mengapa smoking_unknown tercatat 1.483 padahal total Unknown 1.544? Jelaskan pengaruh urutan pemeriksaan.",
        ],
    )

    # ---- DAFTAR PUSTAKA ----
    doc.add_page_break()
    heading1(doc, "Daftar Pustaka")
    for entri in [
        "Slide Pertemuan 3 Machine Learning: Pengkondisian dan Perulangan, S1 Teknologi Informasi.",
        "Fedesoriano. Stroke Prediction Dataset. Kaggle. https://www.kaggle.com/datasets/fedesoriano/stroke-prediction-dataset",
        "Notebook pertemuan3_pengkodisian_perulangan.ipynb, Kelompok 3, repositori machine-learning.",
    ]:
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(entri)
        r.font.name = TNR
        r.font.size = Pt(12)

    doc.save(OUT)
    print("tersimpan:", OUT)


if __name__ == "__main__":
    main()
